# -*- coding: utf-8 -*-
"""
Script to generate report.docx identical to the LaTeX / PDF report
formatted strictly according to SEEE / HUST 2023 Guidelines.
"""

import os
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="8" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="8" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_shading(cell, color_hex="F2F2F2"):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def add_page_number_field(run):
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def create_report_docx(output_path):
    doc = docx.Document()
    
    # ----------------------------------------------------
    # PAGE SETUP: A4, Margins: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm
    # ----------------------------------------------------
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    
    # Normal Style Configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    # Helpers
    def add_para(text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, 
                 space_before=0, space_after=6, line_spacing=1.3, first_indent=1.0, font_size=13):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if first_indent > 0:
            p.paragraph_format.first_line_indent = Cm(first_indent)
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            run.bold = bold
            run.italic = italic
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(14)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text.upper())
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.bold = True
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.bold = True
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.bold = False
        run.italic = True
        return p

    def add_table_caption(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.bold = True
        return p

    def add_figure_caption(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.italic = True
        return p

    # ----------------------------------------------------
    # COVER PAGE (SECTION 0)
    # ----------------------------------------------------
    # Set cover page margins (centered 2cm all around)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    
    # Outer double border frame table
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.autofit = False
    cover_table.columns[0].width = Cm(16.5)
    
    cell = cover_table.cell(0, 0)
    set_cell_margins(cell, top=300, bottom=300, left=350, right=350)
    
    # Set double border on cell
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="double" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="double" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="double" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="double" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    # Content inside cover
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(10)
    cp.paragraph_format.space_after = Pt(2)
    cp.paragraph_format.line_spacing = 1.2
    cp.paragraph_format.first_line_indent = Cm(0)
    r = cp.add_run("ĐẠI HỌC BÁCH KHOA HÀ NỘI\nTRƯỜNG ĐIỆN – ĐIỆN TỬ")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(17)
    r.bold = True
    
    cp_line = cell.add_paragraph()
    cp_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_line.paragraph_format.space_before = Pt(2)
    cp_line.paragraph_format.space_after = Pt(18)
    cp_line.paragraph_format.first_line_indent = Cm(0)
    r_line = cp_line.add_run("—————————————")
    r_line.font.name = 'Times New Roman'
    r_line.font.size = Pt(11)
    r_line.bold = True
    
    # Logo
    cp_logo = cell.add_paragraph()
    cp_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_logo.paragraph_format.space_before = Pt(0)
    cp_logo.paragraph_format.space_after = Pt(20)
    cp_logo.paragraph_format.first_line_indent = Cm(0)
    if os.path.exists("report/src/img/logo.png"):
        cp_logo.add_run().add_picture("report/src/img/logo.png", height=Cm(4.2))
    
    # Title
    cp_title = cell.add_paragraph()
    cp_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_title.paragraph_format.space_before = Pt(10)
    cp_title.paragraph_format.space_after = Pt(8)
    cp_title.paragraph_format.line_spacing = 1.2
    cp_title.paragraph_format.first_line_indent = Cm(0)
    r_title = cp_title.add_run("BÁO CÁO\nTHỰC TẬP KỸ THUẬT")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(20)
    r_title.bold = True
    
    cp_topic = cell.add_paragraph()
    cp_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_topic.paragraph_format.space_before = Pt(10)
    cp_topic.paragraph_format.space_after = Pt(40)
    cp_topic.paragraph_format.line_spacing = 1.2
    cp_topic.paragraph_format.first_line_indent = Cm(0)
    r_topic = cp_topic.add_run("CÔNG TY CỔ PHẦN TECHPAL")
    r_topic.font.name = 'Times New Roman'
    r_topic.font.size = Pt(18)
    r_topic.bold = True
    
    # Info table inside cover
    info_tbl = cell.add_table(rows=4, cols=2)
    info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_tbl.autofit = False
    info_tbl.columns[0].width = Cm(5.2)
    info_tbl.columns[1].width = Cm(8.0)
    
    info_data = [
        ("Giảng viên hướng dẫn:", "TS. Nguyễn Tiến Đạt"),
        ("Sinh viên thực hiện:", "Mai Duy Hiếu"),
        ("MSSV:", "20223967"),
        ("Mã lớp học:", "767288")
    ]
    for row_idx, (k, v) in enumerate(info_data):
        c0 = info_tbl.cell(row_idx, 0)
        c1 = info_tbl.cell(row_idx, 1)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.first_line_indent = Cm(0)
        r0 = p0.add_run(k)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(13)
        r0.bold = True
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.first_line_indent = Cm(0)
        r1 = p1.add_run(v)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(13)
        
    cp_foot = cell.add_paragraph()
    cp_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_foot.paragraph_format.space_before = Pt(50)
    cp_foot.paragraph_format.space_after = Pt(10)
    cp_foot.paragraph_format.first_line_indent = Cm(0)
    r_foot = cp_foot.add_run("Hà Nội, 8/2026")
    r_foot.font.name = 'Times New Roman'
    r_foot.font.size = Pt(13)
    r_foot.bold = True

    # ----------------------------------------------------
    # SECTION 2: FRONTMATTER (Roman page numbering: i, ii, iii...)
    # ----------------------------------------------------
    sec_front = doc.add_section(docx.enum.section.WD_SECTION_START.NEW_PAGE)
    sec_front.top_margin = Cm(2.0)
    sec_front.bottom_margin = Cm(2.0)
    sec_front.left_margin = Cm(3.0)
    sec_front.right_margin = Cm(2.0)
    
    # Configure Roman numerals in footer
    footer_front = sec_front.footer
    footer_p = footer_front.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.first_line_indent = Cm(0)
    footer_run = footer_p.add_run()
    footer_run.font.name = 'Times New Roman'
    footer_run.font.size = Pt(12)
    add_page_number_field(footer_run)
    
    # Set Roman format on section page number
    sectPr = sec_front._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), 'roman')
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)
    
    # --- LỜI NÓI ĐẦU ---
    add_heading_1("LỜI NÓI ĐẦU")
    add_para(
        "Thực tập kỹ thuật là học phần giúp sinh viên kết nối kiến thức cơ sở ngành với môi trường làm việc thực tế. "
        "Đối với ngành Kỹ thuật Điện tử Viễn thông, việc tiếp cận linh kiện điện tử, thiết bị tự động hóa, tài liệu kỹ thuật và quy trình lựa chọn giải pháp có ý nghĩa quan trọng trong quá trình hình thành tư duy kỹ sư. "
        "Đây cũng là cơ hội để sinh viên rèn luyện tác phong làm việc, kỹ năng tìm kiếm thông tin, đọc tài liệu chuyên ngành và trình bày kết quả kỹ thuật."
    )
    add_para(
        "Trong học kỳ 2025.3 – năm học 2025–2026, em thực hiện học phần Thực tập Kỹ thuật tại Công ty Cổ phần TECHPAL, "
        "dưới sự hướng dẫn tại đơn vị của anh Nguyễn Hồng Sơn và sự hướng dẫn của giảng viên TS. Nguyễn Tiến Đạt. "
        "Báo cáo tập trung giới thiệu khái quát doanh nghiệp, các lĩnh vực hoạt động có liên quan đến chuyên ngành, "
        "đồng thời hệ thống hóa nội dung tìm hiểu về linh kiện điện tử và thiết bị tự động hóa."
    )
    add_para(
        "Thuận lợi lớn nhất là Công ty Cổ phần TECHPAL có danh mục linh kiện điện tử và thiết bị tự động hóa khá đa dạng, "
        "kèm trang giới thiệu và trang bán hàng công khai. Nhờ đó em có thể đối chiếu mô tả sản phẩm với datasheet của hãng, "
        "đồng thời được anh Nguyễn Hồng Sơn và giảng viên TS. Nguyễn Tiến Đạt định hướng phạm vi tìm hiểu cho sát chuyên ngành. "
        "Khó khăn chủ yếu đến từ thời gian thực tập ngắn, danh mục mã hàng lớn và cách đặt mã khác nhau giữa các hãng. "
        "Em cũng không tiếp cận được quy trình nội bộ, dữ liệu kho hay phòng thí nghiệm để đo kiểm, nên phần phân tích ứng dụng trong báo cáo mới dừng ở mức phương án trên giấy. "
        "Việc đọc tài liệu tiếng Anh và không bỏ sót hậu tố, kiểu vỏ, dải nhiệt độ hay trạng thái vòng đời cũng đòi hỏi em phải tra cứu nhiều lần mới quen."
    )
    add_para(
        "Em xin chân thành cảm ơn Ban lãnh đạo Công ty Cổ phần TECHPAL, anh Nguyễn Hồng Sơn, giảng viên TS. Nguyễn Tiến Đạt "
        "cùng các cá nhân đã hỗ trợ em trong quá trình thực hiện học phần."
    )
    
    # --- MỤC LỤC ---
    doc.add_page_break()
    add_heading_1("MỤC LỤC")
    
    toc_items = [
        ("LỜI NÓI ĐẦU", "i", True),
        ("DANH MỤC HÌNH VẼ", "iii", True),
        ("DANH MỤC BẢNG", "iv", True),
        ("CHƯƠNG 1. GIỚI THIỆU ĐƠN VỊ THỰC TẬP", "1", True),
        ("   1.1 Thông tin chung", "1", False),
        ("   1.2 Quá trình hình thành và định hướng phát triển", "1", False),
        ("   1.3 Lĩnh vực hoạt động chuyên môn", "2", False),
        ("   1.4 Cơ cấu tổ chức", "3", False),
        ("   1.5 Các vị trí công việc và yêu cầu năng lực", "3", False),
        ("CHƯƠNG 2. NỘI DUNG THỰC TẬP", "4", True),
        ("   2.1 Phạm vi và phương pháp tìm hiểu", "4", False),
        ("   2.2 Các vị trí liên quan trong chuỗi cung ứng", "4", False),
        ("   2.3 Khảo sát danh mục sản phẩm của TECHPAL", "5", False),
        ("   2.4 Quy trình đọc và lựa chọn thiết bị", "7", False),
        ("   2.5 Phân tích một số mã sản phẩm đại diện", "8", False),
        ("   2.6 Bài toán khảo sát: thu thập và truyền dữ liệu nhiệt độ", "10", False),
        ("   2.7 Kết quả và kỹ năng thu nhận", "11", False),
        ("CHƯƠNG 3. NHẬN XÉT VÀ ĐỀ XUẤT", "12", True),
        ("   3.1 Ưu điểm", "12", False),
        ("   3.2 Nhược điểm", "12", False),
        ("   3.3 Đề xuất", "13", False),
        ("KẾT LUẬN", "14", True),
        ("TÀI LIỆU THAM KHẢO", "15", True),
        ("PHỤ LỤC A. THÔNG TIN PHỤC VỤ PHIẾU ĐÁNH GIÁ", "16", True),
        ("   A.1 Thông tin sinh viên và đợt thực tập", "16", False),
        ("   A.2 Nội dung thực tập tóm tắt", "16", False),
        ("PHỤ LỤC B. MINH CHỨNG THỰC TẬP", "17", True),
        ("   B.1 Phiếu đánh giá quá trình thực tập", "17", False),
    ]
    
    toc_tbl = doc.add_table(rows=len(toc_items), cols=2)
    toc_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc_tbl.autofit = False
    toc_tbl.columns[0].width = Cm(14.5)
    toc_tbl.columns[1].width = Cm(1.5)
    
    for idx, (title, pnum, is_bold) in enumerate(toc_items):
        c0 = toc_tbl.cell(idx, 0)
        c1 = toc_tbl.cell(idx, 1)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        p0.paragraph_format.line_spacing = 1.15
        p0.paragraph_format.first_line_indent = Cm(0)
        r0 = p0.add_run(title)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(12)
        r0.bold = is_bold
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        p1.paragraph_format.first_line_indent = Cm(0)
        r1 = p1.add_run(pnum)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)
        r1.bold = is_bold

    # --- DANH MỤC HÌNH VẼ ---
    doc.add_page_break()
    add_heading_1("DANH MỤC HÌNH VẼ")
    
    lof_items = [
        ("Hình 2.1", "Luồng xử lý yêu cầu cung ứng (mô hình hóa từ thông tin công khai)", "5"),
        ("Hình 2.2", "Kiến trúc khái niệm của hệ thống thu thập nhiệt độ", "10"),
        ("Hình B.1", "Phiếu đánh giá quá trình thực tập (trang 1)", "17"),
        ("Hình B.2", "Phiếu đánh giá quá trình thực tập (trang 2)", "18"),
    ]
    lof_tbl = doc.add_table(rows=len(lof_items), cols=3)
    lof_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    lof_tbl.autofit = False
    lof_tbl.columns[0].width = Cm(2.2)
    lof_tbl.columns[1].width = Cm(12.3)
    lof_tbl.columns[2].width = Cm(1.5)
    
    for idx, (h_num, h_title, pnum) in enumerate(lof_items):
        c0 = lof_tbl.cell(idx, 0)
        c1 = lof_tbl.cell(idx, 1)
        c2 = lof_tbl.cell(idx, 2)
        for c, text, align, bold in [(c0, h_num, WD_ALIGN_PARAGRAPH.LEFT, True),
                                     (c1, h_title, WD_ALIGN_PARAGRAPH.LEFT, False),
                                     (c2, pnum, WD_ALIGN_PARAGRAPH.RIGHT, False)]:
            p = c.paragraphs[0]
            p.alignment = align
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.bold = bold

    # --- DANH MỤC BẢNG ---
    doc.add_page_break()
    add_heading_1("DANH MỤC BẢNG")
    
    lot_items = [
        ("Bảng 1.1", "Thông tin khái quát về đơn vị thực tập", "1"),
        ("Bảng 1.2", "Một số nhóm vị trí và năng lực cần thiết", "3"),
        ("Bảng 2.1", "Vai trò của các bộ phận trong chuỗi cung ứng linh kiện", "5"),
        ("Bảng 2.2", "Các nhóm cung ứng chính trên danh mục TECHPAL", "6"),
        ("Bảng 2.3", "Một số mã hàng đại diện trong danh mục TECHPAL", "6"),
        ("Bảng 2.4", "Bảng kiểm kỹ thuật trước khi lựa chọn mã hàng", "8"),
        ("Bảng 2.5", "Yêu cầu đầu vào của bài toán khảo sát", "10"),
        ("Bảng 2.6", "Đánh giá vai trò của các mã sản phẩm trong phương án", "11"),
    ]
    lot_tbl = doc.add_table(rows=len(lot_items), cols=3)
    lot_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    lot_tbl.autofit = False
    lot_tbl.columns[0].width = Cm(2.2)
    lot_tbl.columns[1].width = Cm(12.3)
    lot_tbl.columns[2].width = Cm(1.5)
    
    for idx, (b_num, b_title, pnum) in enumerate(lot_items):
        c0 = lot_tbl.cell(idx, 0)
        c1 = lot_tbl.cell(idx, 1)
        c2 = lot_tbl.cell(idx, 2)
        for c, text, align, bold in [(c0, b_num, WD_ALIGN_PARAGRAPH.LEFT, True),
                                     (c1, b_title, WD_ALIGN_PARAGRAPH.LEFT, False),
                                     (c2, pnum, WD_ALIGN_PARAGRAPH.RIGHT, False)]:
            p = c.paragraphs[0]
            p.alignment = align
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.bold = bold

    # ----------------------------------------------------
    # SECTION 3: MAIN BODY (Arabic page numbering: 1, 2, 3...)
    # ----------------------------------------------------
    sec_body = doc.add_section(docx.enum.section.WD_SECTION_START.NEW_PAGE)
    sec_body.top_margin = Cm(2.0)
    sec_body.bottom_margin = Cm(2.0)
    sec_body.left_margin = Cm(3.0)
    sec_body.right_margin = Cm(2.0)
    
    footer_body = sec_body.footer
    footer_bp = footer_body.paragraphs[0]
    footer_bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_bp.paragraph_format.first_line_indent = Cm(0)
    footer_brun = footer_bp.add_run()
    footer_brun.font.name = 'Times New Roman'
    footer_brun.font.size = Pt(12)
    add_page_number_field(footer_brun)
    
    sectPr_body = sec_body._sectPr
    pgNumType_b = OxmlElement('w:pgNumType')
    pgNumType_b.set(qn('w:fmt'), 'decimal')
    pgNumType_b.set(qn('w:start'), '1')
    sectPr_body.append(pgNumType_b)

    # ====================================================
    # CHƯƠNG 1: GIỚI THIỆU ĐƠN VỊ THỰC TẬP
    # ====================================================
    add_heading_1("CHƯƠNG 1: GIỚI THIỆU ĐƠN VỊ THỰC TẬP")
    
    add_heading_2("1.1 Thông tin chung")
    add_para(
        "Công ty Cổ phần TECHPAL (tên quốc tế: TECHPAL Joint Stock Company, viết tắt TECHPAL., JSC) có mã số thuế 0104469357, thành lập từ ngày 08/02/2010. "
        "Địa chỉ văn phòng giao dịch đặt tại D48-13 Khu đô thị Geleximco, đường Lê Trọng Tấn, phường Dương Nội, Hà Nội. "
        "Số điện thoại liên hệ là 024 3793 8272, email info@techpal.vn, website https://techpal.vn [1]."
    )
    
    add_table_caption("Bảng 1.1: Thông tin khái quát về đơn vị thực tập")
    t1 = doc.add_table(rows=10, cols=2)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.autofit = False
    t1.columns[0].width = Cm(4.5)
    t1.columns[1].width = Cm(11.5)
    set_table_borders(t1)
    
    t1_data = [
        ("Tên doanh nghiệp", "Công ty Cổ phần TECHPAL"),
        ("Tên quốc tế", "TECHPAL Joint Stock Company"),
        ("Tên viết tắt", "TECHPAL., JSC"),
        ("Mã số thuế", "0104469357"),
        ("Ngày thành lập", "08/02/2010"),
        ("Trụ sở chính", "D48-13 Khu đô thị Geleximco, đường Lê Trọng Tấn, phường Dương Nội, Hà Nội"),
        ("Website", "https://techpal.vn"),
        ("Người hướng dẫn", "Nguyễn Hồng Sơn"),
        ("Đại diện xác nhận", "Nguyễn Thế Tài, Phó Tổng giám đốc"),
        ("Thời gian thực tập", "20/07/2026 – 28/08/2026")
    ]
    for r_idx, (k, v) in enumerate(t1_data):
        c0 = t1.cell(r_idx, 0)
        c1 = t1.cell(r_idx, 1)
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.first_line_indent = Cm(0)
        r0 = p0.add_run(k)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(12)
        r0.bold = True
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.first_line_indent = Cm(0)
        r1 = p1.add_run(v)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)

    add_heading_2("1.2 Quá trình hình thành và định hướng phát triển")
    add_para(
        "Trên website chính thức, TECHPAL giới thiệu bốn mảng hoạt động chính gồm: kinh doanh linh kiện điện tử và thiết bị tự động hóa; "
        "nghiên cứu, chuyển giao công nghệ; năng lượng xanh (chủ yếu là điện mặt trời); và đầu tư nông nghiệp công nghệ cao [1]. "
        "Nhìn chung, công ty xuất phát từ lĩnh vực điện – điện tử rồi mở rộng thêm sang năng lượng và nông nghiệp."
    )
    add_para(
        "Về định hướng, TECHPAL tập trung vào đổi mới công nghệ và phát triển bền vững. "
        "Công ty đã triển khai một số dự án điện mặt trời áp mái, song song với hoạt động cung ứng thiết bị tự động hóa và các dự án nông nghiệp ứng dụng công nghệ."
    )

    add_heading_2("1.3 Lĩnh vực hoạt động chuyên môn")
    
    add_heading_3("1.3.1 Linh kiện điện tử và thiết bị tự động hóa")
    add_para(
        "Trên trang bán hàng của TECHPAL [2], sản phẩm được chia thành bốn nhóm: linh kiện điện tử, thiết bị tự động hóa công nghiệp, "
        "vật tư phụ trợ và thiết bị phòng sạch. Em thấy trên danh mục có khá nhiều mã linh kiện bán dẫn, đầu nối, rơ-le, nguồn cấp, cảm biến, "
        "bộ điều khiển PLC, mô-đun vào/ra và bộ ghi dữ liệu."
    )
    add_para(
        "Việc cung ứng linh kiện kỹ thuật không đơn giản chỉ là bán hàng. Người làm kỹ thuật phải đọc đúng mã linh kiện, so sánh thông số điện và cơ khí, "
        "kiểm tra xem có thay thế được không, xem sản phẩm còn được sản xuất hay đã ngừng, rồi cả điều kiện lưu kho, vận chuyển và chống tĩnh điện nữa. "
        "Bỏ sót bước nào cũng có thể gây ra lỗi cho hệ thống về sau."
    )

    add_heading_3("1.3.2 Nghiên cứu và chuyển giao công nghệ")
    add_para(
        "Mảng R&D của TECHPAL hoạt động như cầu nối giữa yêu cầu thực tế và giải pháp kỹ thuật. "
        "Quy trình thường bắt đầu từ khảo sát nhu cầu, sau đó phân rã chức năng, chọn phần cứng, thiết kế kết nối, thử nghiệm rồi bàn giao. "
        "Với sinh viên ngành Điện tử Viễn thông, đây là nơi có thể áp dụng các kiến thức về mạch điện, vi điều khiển, cảm biến hay truyền thông công nghiệp."
    )

    add_heading_3("1.3.3 Năng lượng xanh và điện mặt trời")
    add_para(
        "Một hệ thống điện mặt trời thông thường gồm tấm pin, thiết bị bảo vệ DC, bộ biến đổi (inverter), hệ thống giám sát và phần kết nối lưới. "
        "Khi chọn thiết bị cho hệ thống này cần xét đến công suất, dải điện áp, hiệu suất, điều kiện thời tiết và yêu cầu an toàn. "
        "TECHPAL có tham gia lĩnh vực này thông qua các dự án điện mặt trời áp mái."
    )

    add_heading_3("1.3.4 Nông nghiệp công nghệ cao")
    add_para(
        "Mảng nông nghiệp công nghệ cao tạo ra nhu cầu về cảm biến môi trường, bộ điều khiển, thiết bị truyền thông và lưu trữ dữ liệu. "
        "Tuy mục tiêu cuối cùng là phục vụ sản xuất nông nghiệp, nhưng nhiều thiết bị cốt lõi vẫn thuộc lĩnh vực điện tử và tự động hóa – ví dụ như đo nhiệt độ, độ ẩm, điều khiển quạt, bơm hay thu thập dữ liệu từ xa."
    )

    add_heading_2("1.4 Cơ cấu tổ chức")
    add_para(
        "Theo thông tin trên website [1], TECHPAL có cơ cấu gồm Đại hội đồng cổ đông, Hội đồng quản trị, Ban Kiểm soát, Ban Tổng giám đốc và các phòng chuyên môn. "
        "Cụ thể, các phòng ban được liệt kê gồm: Kinh doanh linh kiện điện tử, Tài chính – Kế toán, Hành chính – Tổng hợp, Phát triển và vận hành dự án, Pháp chế doanh nghiệp. "
        "Ngoài ra còn có đơn vị thành viên, văn phòng đại diện và các ban quản lý dự án. "
        "Bài đăng về Đại hội đồng cổ đông thường niên năm 2026 [10] cũng xác nhận ông Nguyễn Thế Tài giữ chức Thành viên Hội đồng quản trị kiêm Phó Tổng giám đốc."
    )

    add_heading_2("1.5 Các vị trí công việc và yêu cầu năng lực")
    add_para(
        "Từ các lĩnh vực hoạt động nêu trên, em tạm chia ra một số nhóm vị trí liên quan đến chuyên ngành Điện tử Viễn thông như trong Bảng 1.2. "
        "Cần lưu ý đây là cách em tự phân nhóm dựa trên tính chất công việc, không phải danh sách tuyển dụng chính thức của công ty."
    )
    
    add_table_caption("Bảng 1.2: Một số nhóm vị trí và năng lực cần thiết")
    t2 = doc.add_table(rows=6, cols=3)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False
    t2.columns[0].width = Cm(4.0)
    t2.columns[1].width = Cm(6.0)
    t2.columns[2].width = Cm(6.0)
    set_table_borders(t2)
    
    t2_headers = ["Nhóm vị trí", "Nhiệm vụ điển hình", "Kiến thức, kỹ năng cần có"]
    for c_idx, h in enumerate(t2_headers):
        c = t2.cell(0, c_idx)
        set_cell_shading(c, "EAEAEA")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        
    t2_rows = [
        ("Kỹ sư ứng dụng linh kiện", "Đọc yêu cầu, chọn linh kiện, kiểm tra khả năng thay thế và hỗ trợ kỹ thuật", "Mạch điện tử, đọc datasheet, tra cứu mã hàng, giao tiếp kỹ thuật"),
        ("Kỹ sư tự động hóa", "Lựa chọn PLC, cảm biến, cơ cấu chấp hành; cấu hình và kiểm thử hệ thống", "Điều khiển, điện công nghiệp, PLC/HMI, truyền thông công nghiệp"),
        ("Kỹ sư dự án", "Khảo sát, lập phương án, bóc tách thiết bị, theo dõi triển khai", "Tư duy hệ thống, đọc bản vẽ, lập kế hoạch, quản lý rủi ro"),
        ("Kỹ sư R&D", "Phân tích yêu cầu, thiết kế nguyên mẫu, đo kiểm và lập tài liệu", "Điện tử tương tự – số, vi điều khiển, đo lường, thiết kế và kiểm thử"),
        ("Kiểm tra kỹ thuật", "Kiểm tra mã hàng, ngoại quan, thông số và hồ sơ chất lượng", "Dụng cụ đo, tiêu chuẩn chất lượng, ESD, kỹ năng ghi nhận dữ liệu")
    ]
    for r_idx, r_data in enumerate(t2_rows, start=1):
        for c_idx, val in enumerate(r_data):
            c = t2.cell(r_idx, c_idx)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            if c_idx == 0:
                r.bold = True

    # ====================================================
    # CHƯƠNG 2: NỘI DUNG THỰC TẬP
    # ====================================================
    doc.add_page_break()
    add_heading_1("CHƯƠNG 2: NỘI DUNG THỰC TẬP")
    
    add_heading_2("2.1 Phạm vi và phương pháp tìm hiểu")
    add_para(
        "Báo cáo này tập trung vào hoạt động cung ứng linh kiện điện tử và thiết bị tự động hóa tại Công ty Cổ phần TECHPAL. "
        "Em chọn đi vào ba vấn đề chính: thứ nhất là mối liên hệ giữa các bộ phận tham gia chuỗi cung ứng; "
        "thứ hai là cách chuyển yêu cầu của khách hàng thành tiêu chí chọn thiết bị; "
        "và thứ ba là cách kiểm tra một mã hàng cụ thể bằng tài liệu của nhà sản xuất."
    )
    add_para(
        "Nguồn thông tin chủ yếu là website giới thiệu [1], trang bán hàng công khai của TECHPAL [2] và tài liệu kỹ thuật (datasheet) từ các hãng sản xuất. "
        "Em chỉ sử dụng những gì có thể tra cứu và kiểm chứng được, còn quy trình nội bộ, dữ liệu kho hay kết quả đo kiểm thì em không có quyền truy cập. "
        "Các phân tích ứng dụng trong chương này mới ở mức phương án trên giấy, chưa phải sản phẩm đã làm ra."
    )
    add_para("Em thực hiện theo bốn bước:")
    
    steps = [
        "1) Tìm các nhóm sản phẩm và bộ phận liên quan từ nguồn công khai;",
        "2) Chọn vài mã đại diện có trong danh mục của TECHPAL;",
        "3) Đối chiếu mô tả trên trang bán hàng với tài liệu chính thức của hãng;",
        "4) Tổng hợp lại thành quy trình lựa chọn, kèm một bài toán minh họa."
    ]
    for s in steps:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        r = p.add_run(s)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)

    add_heading_2("2.2 Các vị trí liên quan trong chuỗi cung ứng")
    add_para(
        "Trên trang giới thiệu của TECHPAL [3] có đề cập đến ba bộ phận trực tiếp tham gia cung ứng: Sale, Mua hàng Quốc tế và Logistics. "
        "Ba bộ phận này phụ trách các giai đoạn khác nhau nhưng đều làm việc trên cùng một đơn hàng, cùng mã linh kiện, số lượng, yêu cầu chất lượng và deadline giao hàng."
    )
    
    add_table_caption("Bảng 2.1: Vai trò của các bộ phận trong chuỗi cung ứng linh kiện")
    t3 = doc.add_table(rows=4, cols=3)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    t3.autofit = False
    t3.columns[0].width = Cm(3.5)
    t3.columns[1].width = Cm(6.5)
    t3.columns[2].width = Cm(6.0)
    set_table_borders(t3)
    
    t3_headers = ["Bộ phận", "Nội dung công việc liên quan", "Kiến thức, kỹ năng cần thiết"]
    for c_idx, h in enumerate(t3_headers):
        c = t3.cell(0, c_idx)
        set_cell_shading(c, "EAEAEA")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        
    t3_rows = [
        ("Sale", "Tiếp nhận nhu cầu, làm rõ mã hàng, số lượng, ứng dụng và thời hạn; trao đổi phương án với khách hàng", "Giao tiếp kỹ thuật, đọc mã linh kiện, hiểu thông số cơ bản và lập báo giá"),
        ("Mua hàng Quốc tế", "Tìm nguồn, kiểm tra khả năng cung ứng, quy cách đóng gói và chứng từ; làm việc với nhà cung cấp", "Tiếng Anh chuyên ngành, tra cứu chéo, đánh giá nguồn và quản lý thay đổi sản phẩm"),
        ("Logistics", "Theo dõi vận chuyển, tiếp nhận, bảo quản và bàn giao; kiểm tra sự phù hợp giữa hàng hóa với chứng từ", "Quản lý mã hàng, đóng gói, điều kiện lưu kho, ESD và truy xuất lô hàng")
    ]
    for r_idx, r_data in enumerate(t3_rows, start=1):
        for c_idx, val in enumerate(r_data):
            c = t3.cell(r_idx, c_idx)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            if c_idx == 0:
                r.bold = True

    add_para(
        "Bảng 2.1 là do em tự tổng hợp theo tính chất công việc, không phải bản mô tả nhân sự chính thức của công ty. "
        "Tuy nhiên, một điểm em thấy rõ là với linh kiện điện tử, chỉ cần sai một hậu tố trong mã đặt hàng là có thể nhận nhầm kiểu vỏ, dải nhiệt độ hay quy cách đóng gói. "
        "Cho nên thông tin kỹ thuật cần phải được giữ nhất quán khi chuyển từ bộ phận này sang bộ phận khác."
    )
    add_para(
        "Từ thông tin về dịch vụ đặt hàng qua châu Âu – Mỹ, DigiKey, Mouser và Trung Quốc trên trang TECHPAL [3], "
        "em vẽ lại luồng xử lý một yêu cầu cung ứng như Hình 2.1. Sơ đồ này chỉ mang tính khái quát để phục vụ học tập, không phải quy trình nội bộ mà công ty đang áp dụng."
    )
    
    # Diagram 2.1
    p_fig21 = doc.add_paragraph()
    p_fig21.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fig21.paragraph_format.space_before = Pt(6)
    p_fig21.paragraph_format.space_after = Pt(2)
    p_fig21.paragraph_format.first_line_indent = Cm(0)
    if os.path.exists("report/src/img/fig_supply_flow.png"):
        p_fig21.add_run().add_picture("report/src/img/fig_supply_flow.png", width=Cm(15.5))
    add_figure_caption("Hình 2.1: Luồng xử lý yêu cầu cung ứng (mô hình hóa từ thông tin công khai)")

    add_heading_2("2.3 Khảo sát danh mục sản phẩm của TECHPAL")
    add_para(
        "Danh mục trực tuyến của TECHPAL [2] chia sản phẩm thành bốn nhóm lớn: linh kiện điện tử; thiết bị tự động hóa công nghiệp; "
        "vật tư phụ trợ công nghiệp; và thiết bị, vật tư phòng sạch. Cách phân nhóm này giúp thu hẹp phạm vi tìm kiếm, "
        "nhưng chỉ nhìn vào tên nhóm thì chưa biết được sản phẩm có hợp với ứng dụng cụ thể hay không."
    )
    
    add_table_caption("Bảng 2.2: Các nhóm cung ứng chính trên danh mục TECHPAL")
    t4 = doc.add_table(rows=5, cols=3)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    t4.autofit = False
    t4.columns[0].width = Cm(4.0)
    t4.columns[1].width = Cm(6.0)
    t4.columns[2].width = Cm(6.0)
    set_table_borders(t4)
    
    t4_headers = ["Nhóm cung ứng", "Thành phần tiêu biểu quan sát được", "Thông tin cần kiểm tra"]
    for c_idx, h in enumerate(t4_headers):
        c = t4.cell(0, c_idx)
        set_cell_shading(c, "EAEAEA")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        
    t4_rows = [
        ("Linh kiện điện tử", "IC, vi điều khiển, linh kiện công suất, đầu nối, rơ-le, màn hình", "Mã đầy đủ, kiểu vỏ, điện áp, dòng, nhiệt độ và trạng thái vòng đời"),
        ("Thiết bị tự động hóa", "Mô-đun PLC/I/O, cảm biến, encoder, bộ nguồn, thiết bị ghi dữ liệu", "Kiến trúc hệ thống, số kênh, loại tín hiệu, giao thức và khả năng mở rộng"),
        ("Vật tư phụ trợ công nghiệp", "Đầu nối, dây dẫn, phụ kiện lắp đặt và vật tư bảo trì", "Kích thước, vật liệu, điều kiện môi trường và tương thích cơ khí"),
        ("Thiết bị, vật tư phòng sạch", "Dụng cụ và vật tư phục vụ môi trường kiểm soát bụi bẩn", "Cấp sạch, vật liệu, quy cách đóng gói và yêu cầu sử dụng")
    ]
    for r_idx, r_data in enumerate(t4_rows, start=1):
        for c_idx, val in enumerate(r_data):
            c = t4.cell(r_idx, c_idx)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            if c_idx == 0:
                r.bold = True

    add_para(
        "Để không dừng lại ở mô tả chung chung, em chọn bốn mã hàng có thật trong danh mục TECHPAL để tra cứu kỹ hơn. "
        "Bốn mã này đại diện cho bốn loại thiết bị khác nhau: IC truyền thông, mô-đun I/O công nghiệp, thiết bị kết nối I/O phân tán và bộ ghi dữ liệu."
    )
    
    add_table_caption("Bảng 2.3: Một số mã hàng đại diện trong danh mục TECHPAL")
    t5 = doc.add_table(rows=5, cols=3)
    t5.alignment = WD_TABLE_ALIGNMENT.CENTER
    t5.autofit = False
    t5.columns[0].width = Cm(4.5)
    t5.columns[1].width = Cm(4.0)
    t5.columns[2].width = Cm(7.5)
    set_table_borders(t5)
    
    t5_headers = ["Mã sản phẩm", "Nhóm chức năng", "Mô tả dùng cho bước tra cứu"]
    for c_idx, h in enumerate(t5_headers):
        c = t5.cell(0, c_idx)
        set_cell_shading(c, "EAEAEA")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        
    t5_rows = [
        ("MAX487CSA+T", "IC truyền thông", "Bộ thu phát RS-485/RS-422 công suất thấp, vỏ SOIC 8 chân"),
        ("BMXAMI0410", "Mô-đun vào tương tự", "Mô-đun Modicon X80 có bốn ngõ vào tương tự tốc độ cao"),
        ("6ES7153-1AA03-0XB0", "Giao tiếp I/O phân tán", "Mô-đun kết nối SIMATIC DP IM 153-1 cho ET 200M"),
        ("CR1000X", "Đo lường và lưu trữ", "Bộ ghi dữ liệu có khả năng đo, xử lý, lưu trữ và truyền thông")
    ]
    for r_idx, r_data in enumerate(t5_rows, start=1):
        for c_idx, val in enumerate(r_data):
            c = t5.cell(r_idx, c_idx)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            if c_idx == 0:
                r.font.name = 'Consolas'

    add_heading_2("2.4 Quy trình đọc và lựa chọn thiết bị")
    
    add_heading_3("2.4.1 Chuẩn hóa yêu cầu đầu vào")
    add_para(
        "Nếu khách hàng chỉ nói “cần IC RS-485” hay “cần mô-đun analog” thì chưa đủ thông tin để đặt hàng. "
        "Yêu cầu như vậy cần được cụ thể hóa thành: chức năng gì, điện áp bao nhiêu, mấy kênh, loại tín hiệu, tốc độ, giao thức, nhiệt độ làm việc, "
        "kiểu lắp, số lượng và khi nào cần. Nếu đây là mã thay thế cho linh kiện cũ thì phải có thêm mã gốc và thông tin về bo mạch đang sử dụng."
    )

    add_heading_3("2.4.2 Kiểm tra mã và tài liệu kỹ thuật")
    add_para(
        "Mã hàng cần đọc đầy đủ kể cả hậu tố – điểm này ban đầu em cũng hay bỏ qua. "
        "Để xác nhận thông số, nguồn đáng tin nhất là trang sản phẩm và datasheet của nhà sản xuất. "
        "Mô tả trên website nhà cung ứng giúp định vị nhanh sản phẩm, nhưng không nên dựa vào đó để kết luận tương thích hay không. "
        "Các bước kiểm tra em đề xuất gồm:"
    )
    
    check_steps = [
        "1) Xác nhận nhà sản xuất và mã đặt hàng đầy đủ;",
        "2) Đối chiếu chức năng, sơ đồ chân hoặc kiến trúc mô-đun;",
        "3) Kiểm tra giới hạn điện, nhiệt và điều kiện làm việc khuyến nghị;",
        "4) Xem kiểu vỏ, kích thước, đầu nối và cách lắp đặt;",
        "5) Xác định trạng thái vòng đời: còn sản xuất hay chỉ còn dạng phụ tùng;",
        "6) Lưu lại link tài liệu, ngày tra cứu và kết luận."
    ]
    for cs in check_steps:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        r = p.add_run(cs)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)

    add_para(
        "Kết quả tra cứu nên ghi lại dưới dạng bảng kiểm để sau này người khác có thể xem lại được nguồn dữ liệu và lý do chọn. "
        "Bảng 2.4 liệt kê những câu hỏi tối thiểu cần trả lời trước khi đề xuất mua một mã hàng."
    )
    
    add_table_caption("Bảng 2.4: Bảng kiểm kỹ thuật trước khi lựa chọn mã hàng")
    t6 = doc.add_table(rows=8, cols=3)
    t6.alignment = WD_TABLE_ALIGNMENT.CENTER
    t6.autofit = False
    t6.columns[0].width = Cm(3.5)
    t6.columns[1].width = Cm(6.5)
    t6.columns[2].width = Cm(6.0)
    set_table_borders(t6)
    
    t6_headers = ["Nội dung kiểm tra", "Câu hỏi cần trả lời", "Rủi ro nếu bỏ sót"]
    for c_idx, h in enumerate(t6_headers):
        c = t6.cell(0, c_idx)
        set_cell_shading(c, "EAEAEA")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        
    t6_rows = [
        ("Định danh", "Mã, hậu tố và hãng đã chính xác chưa?", "Nhận sai biến thể hoặc sai quy cách đóng gói"),
        ("Thông số điện", "Điện áp, dòng, mức logic và công suất có phù hợp?", "Hỏng linh kiện hoặc hệ thống hoạt động không ổn định"),
        ("Giao tiếp", "Chuẩn vật lý, giao thức và tốc độ có tương thích?", "Không trao đổi được dữ liệu"),
        ("Cơ khí", "Vỏ, chân, đầu nối và kích thước lắp có khớp?", "Không lắp được hoặc phải sửa thiết kế"),
        ("Môi trường", "Nhiệt độ, độ ẩm, rung và cấp bảo vệ có đáp ứng?", "Suy giảm độ tin cậy khi vận hành"),
        ("Vòng đời", "Sản phẩm còn sản xuất hay chỉ còn dạng phụ tùng?", "Khó mua bổ sung và tăng chi phí bảo trì"),
        ("Chứng từ", "Datasheet, nguồn gốc, lô hàng và chứng nhận có đủ?", "Khó truy xuất và nghiệm thu")
    ]
    for r_idx, r_data in enumerate(t6_rows, start=1):
        for c_idx, val in enumerate(r_data):
            c = t6.cell(r_idx, c_idx)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            if c_idx == 0:
                r.bold = True

    add_heading_3("2.4.3 Đánh giá khả năng thay thế")
    add_para(
        "Hai sản phẩm có cùng tên chức năng chưa chắc đã thay thế được cho nhau. "
        "Với IC thì phải so chân, mức điện, tốc độ, tải bus, vỏ và dải nhiệt độ. "
        "Với mô-đun công nghiệp thì còn phải xem rack, bộ điều khiển đi kèm, phiên bản firmware, cấu hình I/O và chứng nhận. "
        "Đặc biệt khi sản phẩm đã ngừng sản xuất (chỉ còn dạng phụ tùng), chi phí thay thế phải tính cả việc chuyển đổi hệ thống chứ không chỉ nhìn giá mô-đun."
    )

    add_heading_2("2.5 Phân tích một số mã sản phẩm đại diện")
    
    add_heading_3("2.5.1 Bộ thu phát MAX487CSA+T")
    add_para(
        "Trên trang TECHPAL [4], MAX487CSA+T được liệt kê trong nhóm linh kiện điện tử, mô tả là bộ thu phát bán song công, vỏ SOIC 8 chân. "
        "Em tra trên trang Analog Devices [5] thì thấy MAX487 là IC thu phát RS-485/RS-422 nguồn đơn 5 V, phiên bản giới hạn tốc độ sườn, hỗ trợ tới 250 kbps. "
        "Đầu vào thu có tải 1/4 đơn vị nên về lý thuyết có thể nối tối đa 128 nút trên một bus."
    )
    add_para(
        "Nếu dùng MAX487CSA+T cho giao tiếp nối tiếp trong thiết kế nhúng thì cần kiểm tra thêm mấy thứ: "
        "điện áp logic phía vi điều khiển có tương thích 5 V không, đường truyền có cần điện trở kết thúc không, bus cần phân cực ra sao, "
        "có cần bảo vệ ESD không, và nếu hai đầu chênh lệch điện thế đất thì có cần cách ly không. "
        "Ngoài ra, hậu tố “T” liên quan tới quy cách đóng gói (tape and reel) nên khi đặt mua không được bỏ."
    )

    add_heading_3("2.5.2 Mô-đun vào tương tự BMXAMI0410")
    add_para(
        "Mã BMXAMI0410 trên TECHPAL thuộc nhóm thiết bị tự động hóa. Theo Schneider Electric [6], đây là mô-đun Modicon X80 với bốn ngõ vào analog tốc độ cao. "
        "Khác với một IC rời, mô-đun này không thể dùng đơn lẻ mà phụ thuộc vào cả hệ thống: rack nào tương thích, CPU PLC nào, "
        "cảm biến ra kiểu tín hiệu gì, đấu dây thế nào, phần mềm cấu hình ra sao."
    )
    add_para(
        "Điều em rút ra ở đây là trong cung ứng tự động hóa, số kênh hay loại tín hiệu mới chỉ là điều kiện cần. "
        "Muốn chốt mã thì phải xem thêm tài liệu tương thích hệ thống và phụ kiện đầu nối. "
        "Giả sử khách hàng chỉ nói “cho mô-đun analog bốn kênh”, bên kỹ thuật hoặc Sale cần hỏi thêm đang dùng PLC/rack gì và đo tín hiệu dải nào."
    )

    add_heading_3("2.5.3 Mô-đun Siemens 6ES7153-1AA03-0XB0")
    add_para(
        "TECHPAL có bán mã 6ES7153-1AA03-0XB0. Tra trên trang Siemens [7], em thấy đây là mô-đun kết nối SIMATIC DP IM 153-1 dùng cho trạm ET 200M, hỗ trợ tối đa tám mô-đun S7-300. "
        "Điểm đáng chú ý là mã này hiện tại Siemens phân loại là phụ tùng (spare part only)."
    )
    add_para(
        "Trạng thái phụ tùng ảnh hưởng rất lớn đến quyết định mua. Nếu nhà máy đang chạy hệ thống cũ và cần thay thế để giảm thời gian dừng máy thì vẫn phải mua mã này. "
        "Nhưng nếu làm dự án mới thì nên cân nhắc nền tảng khác có vòng đời dài hơn. "
        "Em nghĩ đây là ví dụ rõ nhất cho thấy kiểm tra vòng đời quan trọng không kém kiểm tra thông số kỹ thuật."
    )

    add_heading_3("2.5.4 Bộ ghi dữ liệu CR1000X")
    add_para(
        "TECHPAL liệt kê CR1000X của Campbell Scientific trong nhóm thiết bị tự động hóa, mô tả là bộ lưu trữ xử lý dữ liệu [8]. "
        "Tài liệu hãng [9] cho biết CR1000X có 16 đầu vào analog đơn (hoặc 8 vi sai), ADC 24 bit, và tích hợp sẵn khả năng đo cảm biến, xử lý, lưu trữ lẫn truyền dữ liệu."
    )
    add_para(
        "CR1000X khá phù hợp cho các bài toán thu thập dữ liệu môi trường hay trạm đo từ xa. "
        "Tuy vậy, không phải cứ thấy 16 kênh là đủ. Người chọn thiết bị phải xem kỹ: loại cảm biến, dải đo, độ chính xác cần thiết, chu kỳ lấy mẫu, "
        "bộ nhớ, cổng giao tiếp, nguồn cấp và điều kiện ngoài trời hay trong nhà. Hãng cũng khuyến cáo nên xác nhận cấu hình trước khi đặt mua."
    )

    add_heading_2("2.6 Bài toán khảo sát: thu thập và truyền dữ liệu nhiệt độ")
    
    add_heading_3("2.6.1 Yêu cầu bài toán")
    add_para(
        "Để áp dụng thử quy trình đã nêu, em tự đặt ra một bài toán: thiết kế phương án thu thập nhiệt độ trong tủ kỹ thuật. "
        "Hệ thống cần đo được nhiệt độ, lưu giá trị theo thời gian, cảnh báo khi vượt ngưỡng và truyền dữ liệu về máy tính giám sát. "
        "Cần nhấn mạnh rằng đây chỉ là phương án trên giấy, dùng để so sánh các nhóm thiết bị trong danh mục TECHPAL, không phải hệ thống đã lắp đặt hay thử nghiệm."
    )
    
    add_table_caption("Bảng 2.5: Yêu cầu đầu vào của bài toán khảo sát")
    t7 = doc.add_table(rows=7, cols=2)
    t7.alignment = WD_TABLE_ALIGNMENT.CENTER
    t7.autofit = False
    t7.columns[0].width = Cm(4.5)
    t7.columns[1].width = Cm(11.5)
    set_table_borders(t7)
    
    t7_data = [
        ("Đại lượng đo", "Nhiệt độ không khí trong tủ kỹ thuật"),
        ("Số điểm đo", "Từ một điểm, có khả năng mở rộng"),
        ("Chu kỳ lưu", "Có thể cấu hình theo tốc độ biến đổi nhiệt"),
        ("Chức năng", "Hiển thị, lưu dữ liệu và cảnh báo vượt ngưỡng"),
        ("Truyền thông", "Kết nối tới máy tính hoặc mạng giám sát"),
        ("Nguồn", "Chọn theo thiết bị, có bảo vệ và dự phòng phù hợp"),
        ("Giới hạn báo cáo", "Phân tích kiến trúc, chưa chế tạo và chưa có số liệu đo")
    ]
    for r_idx, (k, v) in enumerate(t7_data):
        c0 = t7.cell(r_idx, 0)
        c1 = t7.cell(r_idx, 1)
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.first_line_indent = Cm(0)
        r0 = p0.add_run(k)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(12)
        r0.bold = True
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.first_line_indent = Cm(0)
        r1 = p1.add_run(v)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)

    add_heading_3("2.6.2 Kiến trúc đề xuất")
    add_para(
        "Tùy vào hệ thống hiện có của khách hàng, phần thu thập dữ liệu có thể là bộ ghi chuyên dụng (như CR1000X), "
        "mô-đun analog trong hệ PLC (như BMXAMI0410), hoặc một bo mạch nhúng tự thiết kế dùng IC RS-485 (như MAX487). "
        "Ba hướng này không thay thế nhau trực tiếp mà phù hợp với các mức độ tích hợp khác nhau."
    )
    
    # Diagram 2.2
    p_fig22 = doc.add_paragraph()
    p_fig22.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fig22.paragraph_format.space_before = Pt(6)
    p_fig22.paragraph_format.space_after = Pt(2)
    p_fig22.paragraph_format.first_line_indent = Cm(0)
    if os.path.exists("report/src/img/fig_temperature_architecture.png"):
        p_fig22.add_run().add_picture("report/src/img/fig_temperature_architecture.png", width=Cm(15.5))
    add_figure_caption("Hình 2.2: Kiến trúc khái niệm của hệ thống thu thập nhiệt độ")

    add_table_caption("Bảng 2.6: Đánh giá vai trò của các mã sản phẩm trong phương án")
    t8 = doc.add_table(rows=5, cols=3)
    t8.alignment = WD_TABLE_ALIGNMENT.CENTER
    t8.autofit = False
    t8.columns[0].width = Cm(4.5)
    t8.columns[1].width = Cm(5.5)
    t8.columns[2].width = Cm(6.0)
    set_table_borders(t8)
    
    t8_headers = ["Mã sản phẩm", "Vai trò có thể xem xét", "Điều kiện cần xác minh"]
    for c_idx, h in enumerate(t8_headers):
        c = t8.cell(0, c_idx)
        set_cell_shading(c, "EAEAEA")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        
    t8_rows = [
        ("CR1000X", "Thu thập, xử lý và lưu dữ liệu trong một thiết bị chuyên dụng", "Loại cảm biến, dải đo, chu kỳ lấy mẫu, nguồn và giao tiếp"),
        ("BMXAMI0410", "Nhận tín hiệu analog trong hệ PLC Modicon X80", "Rack, CPU, đầu nối, loại tín hiệu và cấu hình phần mềm"),
        ("MAX487CSA+T", "Tạo giao tiếp RS-485 cho phương án bộ xử lý nhúng", "Nguồn 5 V, mức logic, topology bus, kết thúc và bảo vệ"),
        ("6ES7153-1AA03-0XB0", "Duy trì kết nối I/O phân tán trong hệ ET 200M hiện hữu", "Cấu hình trạm, tương thích S7-300 và trạng thái phụ tùng")
    ]
    for r_idx, r_data in enumerate(t8_rows, start=1):
        for c_idx, val in enumerate(r_data):
            c = t8.cell(r_idx, c_idx)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            if c_idx == 0:
                r.font.name = 'Consolas'

    add_para(
        "Thực tế thì chưa thể chốt phương án nào cho đến khi biết rõ khách hàng dùng cảm biến gì, đang có sẵn nền tảng điều khiển nào và yêu cầu vận hành cụ thể ra sao. "
        "Nếu cần một trạm đo độc lập, bộ ghi dữ liệu kiểu CR1000X sẽ tiện vì nó tích hợp sẵn mọi thứ. "
        "Nếu nhà máy đã dùng PLC Modicon X80 rồi, mô-đun analog tương thích sẽ dễ bảo trì hơn. "
        "Còn nếu tự thiết kế bo mạch nhúng, IC RS-485 chỉ là phần giao tiếp, không thay thế được bộ thu thập dữ liệu."
    )

    add_heading_2("2.7 Kết quả và kỹ năng thu nhận")
    add_para(
        "Sau quá trình tìm hiểu, em đã hệ thống hóa được bốn nhóm cung ứng của TECHPAL, "
        "vẽ được sơ đồ liên hệ giữa Sale, Mua hàng Quốc tế và Logistics, lập một bảng kiểm kỹ thuật và đối chiếu bốn mã sản phẩm với datasheet của hãng. "
        "Bài toán nhiệt độ giúp em thử áp dụng cách chuyển từ yêu cầu chức năng sang kiến trúc và danh sách thông tin cần kiểm tra."
    )
    add_para(
        "Điều em thấy quan trọng nhất là không nên ra quyết định chỉ vì tên sản phẩm nghe hợp lý hay mô tả ngắn trên web trùng khớp. "
        "Mỗi mã linh kiện đều cần kiểm tra bằng datasheet, tài liệu hệ thống và trạng thái vòng đời. "
        "Ngoài kiến thức chuyên ngành, công việc cung ứng còn đòi hỏi khả năng đọc tiếng Anh, ghi chép cẩn thận và phối hợp giữa nhiều bộ phận. "
        "Báo cáo của em dừng ở mức khảo sát tài liệu và đề xuất phương án, chưa có nguyên mẫu hoặc số liệu đo thực tế."
    )

    # ====================================================
    # CHƯƠNG 3: NHẬN XÉT VÀ ĐỀ XUẤT
    # ====================================================
    doc.add_page_break()
    add_heading_1("CHƯƠNG 3: NHẬN XÉT VÀ ĐỀ XUẤT")
    
    add_heading_2("3.1 Ưu điểm")
    add_para(
        "Đề tài bám đúng chuyên ngành Kỹ thuật Điện tử Viễn thông: đọc mã linh kiện, đối chiếu datasheet và xem thiết bị tự động hóa được chọn như thế nào trong một doanh nghiệp cung ứng thật. "
        "Danh mục công khai của TECHPAL đủ rộng để em không phải dừng ở mô tả chung, mà đi vào bốn nhóm sản phẩm và bốn mã đại diện."
    )
    add_para(
        "Việc tra cứu MAX487CSA+T, BMXAMI0410, 6ES7153-1AA03-0XB0 và CR1000X giúp em thấy cùng một bài toán chức năng có thể đi theo các hướng khác nhau. "
        "IC rời phù hợp thiết kế nhúng; mô-đun analog phải khớp rack và CPU; bộ ghi dữ liệu tích hợp tiện cho trạm đo độc lập. "
        "Trường hợp Siemens phân loại 6ES7153-1AA03-0XB0 là phụ tùng là ví dụ rõ nhất: nếu chỉ nhìn tên chức năng thì dễ đề xuất sai cho dự án mới."
    )
    add_para(
        "Em cũng tập được thói quen không tin mô tả ngắn trên trang bán hàng. Bảng kiểm kỹ thuật (Bảng 2.4) buộc phải trả lời đủ định danh, thông số điện, "
        "giao tiếp, cơ khí, môi trường, vòng đời và chứng từ trước khi chốt mã. "
        "Bài toán thu thập nhiệt độ tủ kỹ thuật dù chỉ ở mức phương án giấy vẫn hữu ích, vì nó buộc em đi từ yêu cầu chức năng sang kiến trúc rồi mới so sánh vai trò từng mã. "
        "Các kỹ năng đọc tiếng Anh chuyên ngành, ghi chép nguồn và trình bày phương án là phần em thấy mang sang các học phần thiết kế sau này được."
    )

    add_heading_2("3.2 Nhược điểm")
    add_para(
        "Về điều kiện thực tập, thời gian ngắn nên em chỉ kịp đi sâu một mảng cung ứng. "
        "Em không tiếp cận được quy trình nội bộ, dữ liệu kho hay phòng thí nghiệm, vì vậy Hình 2.1 chỉ là mô hình hóa từ thông tin công khai, không phải luồng đang áp dụng tại công ty. "
        "Bài toán nhiệt độ cũng chưa chọn cảm biến, chưa có số liệu đo, chưa đánh giá được sai số, nhiễu hay độ bền khi vận hành."
    )
    add_para(
        "Về phía bản thân, danh mục công nghiệp lớn và cách đặt mã khác nhau giữa các hãng khiến em dễ chú ý vài thông số nổi (số kênh, điện áp) rồi bỏ sót chi tiết quyết định. "
        "Hậu tố “T” của MAX487CSA+T liên quan quy cách đóng gói tape and reel, không phải biến thể mạch; nếu đặt hàng thiếu hậu tố có thể nhận sai dạng cuộn. "
        "Với BMXAMI0410, em chưa đọc được tài liệu tương thích rack/CPU ở mức triển khai. "
        "Với mã Siemens, nếu không tra trạng thái vòng đời thì sẽ không biết sản phẩm chỉ còn dạng phụ tùng. "
        "Do đó mọi kết luận về khả năng thay thế trong báo cáo đều cần người có kinh nghiệm xem lại."
    )
    add_para(
        "Nhược điểm này cũng thể hiện ở cách viết: một số đoạn còn kể lại thông tin web nhiều hơn là phân tích quyết định kỹ thuật. "
        "Em nhận ra rằng khảo sát tài liệu là bước cần, nhưng chưa đủ để khẳng định một mã “dùng được” trong hệ thống cụ thể."
    )

    add_heading_2("3.3 Đề xuất")
    add_para(
        "Với đơn vị tiếp nhận, em đề xuất có sẵn một bộ tài liệu ngắn cho sinh viên mới: quy tắc đọc mã và hậu tố, "
        "mẫu so sánh thiết bị theo các cột của bảng kiểm, hướng dẫn chống tĩnh điện (ESD), kèm một tình huống chọn mã thật – ví dụ một yêu cầu thay thế hoặc một đơn hàng đã xử lý – "
        "để sinh viên quan sát cách Sale, Mua hàng Quốc tế và Logistics giữ thông tin nhất quán. "
        "Nếu điều kiện cho phép, cho sinh viên xem một lần kiểm tra ngoại quan hoặc đối chiếu chứng từ khi nhận hàng sẽ giúp nối phần đọc datasheet với công việc thực tế, dù không cần giao việc nội bộ."
    )
    add_para(
        "Với bản thân, em sẽ hoàn thiện bảng kiểm cho đủ bốn mã đã chọn, đặc biệt các cột vòng đời và tương thích hệ thống còn mỏng. "
        "Nếu có điều kiện, em muốn làm một nguyên mẫu nhỏ đo nhiệt độ tủ kỹ thuật để kiểm chứng phương án giấy: chọn cảm biến, xác định dải đo, chu kỳ lưu và cổng truyền về máy tính. "
        "Song song đó, em cần luyện viết ghi chú kỹ thuật ngắn hơn, tập trung vào tiêu chí chốt mã thay vì thuật lại toàn bộ trang sản phẩm. "
        "Những việc này sẽ giúp em chắc hơn khi học tiếp thiết kế mạch, hệ thống nhúng và tự động hóa."
    )

    # ====================================================
    # KẾT LUẬN
    # ====================================================
    doc.add_page_break()
    add_heading_1("KẾT LUẬN")
    add_para(
        "Qua đợt thực tập, em đã tìm hiểu được tổng quan về Công ty Cổ phần TECHPAL – từ lịch sử hình thành, lĩnh vực hoạt động cho đến cơ cấu tổ chức. "
        "Phần chính của báo cáo tập trung vào cách đọc tài liệu, tra cứu và lựa chọn linh kiện điện tử, thiết bị tự động hóa dựa trên danh mục sản phẩm của TECHPAL. "
        "Em cũng thử áp dụng kiến thức đó vào một bài toán giám sát nhiệt độ tủ kỹ thuật để minh họa cách đi từ yêu cầu đến sơ đồ khối và so sánh các phương án."
    )
    add_para(
        "Bài học lớn nhất em rút ra là khi chọn linh kiện, không thể chỉ nhìn vào tên gọi hay mô tả ngắn. "
        "Phải kiểm tra đồng thời nhiều yếu tố: giới hạn điện, dải nhiệt độ, kích thước cơ khí, khả năng thay thế, vòng đời sản phẩm và yêu cầu bảo trì. "
        "Cách làm việc có hệ thống như vậy sẽ giúp em học tiếp các môn thiết kế mạch, hệ thống nhúng và tự động hóa tốt hơn."
    )

    # ====================================================
    # TÀI LIỆU THAM KHẢO
    # ====================================================
    doc.add_page_break()
    add_heading_1("TÀI LIỆU THAM KHẢO")
    
    references = [
        "[1] Công ty Cổ phần TECHPAL, “Giới thiệu TECHPAL”, https://www.techpal.com.vn/gioi-thieu/, truy cập tháng 8/2026.",
        "[2] Công ty Cổ phần TECHPAL, “Siêu thị linh kiện điện tử TECHPAL”, https://linhkien.techpal.vn/, truy cập tháng 8/2026.",
        "[3] Công ty Cổ phần TECHPAL, “Giới thiệu Siêu thị linh kiện điện tử TECHPAL”, https://linhkien.techpal.vn/about-us/, truy cập tháng 8/2026.",
        "[4] Công ty Cổ phần TECHPAL, “MAX487CSA+T”, https://linhkien.techpal.vn/san-pham/max487csat/, truy cập tháng 8/2026.",
        "[5] Analog Devices, “MAX487: Low-Power, Slew-Rate-Limited RS-485/RS-422 Transceiver”, https://www.analog.com/en/products/max487.html, truy cập tháng 8/2026.",
        "[6] Schneider Electric, “BMXAMI0410 – Analog Input Module X80, 4 Inputs, High Speed”, https://www.se.com/in/en/product/BMXAMI0410/analog-input-module-x80-4-inputs-high-speed/, truy cập tháng 8/2026.",
        "[7] Siemens, “6ES7153-1AA03-0XB0 – SIMATIC DP IM 153-1”, https://mall.industry.siemens.com/mall/en/oeii/Catalog/Product/6ES7153-1AA03-0XB0, truy cập tháng 8/2026.",
        "[8] Công ty Cổ phần TECHPAL, “Bộ lưu trữ xử lý dữ liệu Datalogger CR1000X”, https://linhkien.techpal.vn/san-pham/bo-luu-tru-xu-ly-du-lieu-datalogger-cr1000x/, truy cập tháng 8/2026.",
        "[9] Campbell Scientific, “CR1000X Specifications”, https://help.campbellsci.com/CR1000X/Content/shared/SharedSpecs/Specifications-cr1000x.htm, truy cập tháng 8/2026.",
        "[10] Công ty Cổ phần TECHPAL, “TECHPAL tổ chức thành công Đại hội đồng cổ đông thường niên năm 2026”, https://www.techpal.vn/techpal-to-chuc-thanh-cong-dai-hoi-dong-co-dong-thuong-nien-nam-2026/, truy cập tháng 8/2026.",
        "[11] Trường Điện – Điện tử, Đại học Bách khoa Hà Nội, “Hướng dẫn viết nội dung báo cáo kết quả Thực tập Kỹ thuật”, 2023."
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        r = p.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # ====================================================
    # PHỤ LỤC A: THÔNG TIN PHỤC VỤ PHIẾU ĐÁNH GIÁ
    # ====================================================
    doc.add_page_break()
    add_heading_1("PHỤ LỤC A: THÔNG TIN PHỤC VỤ PHIẾU ĐÁNH GIÁ")
    
    add_heading_2("A.1 Thông tin sinh viên và đợt thực tập")
    
    t_appA = doc.add_table(rows=10, cols=2)
    t_appA.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_appA.autofit = False
    t_appA.columns[0].width = Cm(4.5)
    t_appA.columns[1].width = Cm(11.5)
    set_table_borders(t_appA)
    
    appA_data = [
        ("Họ và tên", "Mai Duy Hiếu"),
        ("MSSV", "20223967"),
        ("Lớp", "767288"),
        ("Ngày sinh", "12/09/2004"),
        ("Nơi sinh", "Xã Hải Vân, huyện Hải Hậu, tỉnh Nam Định"),
        ("Cơ quan thực tập", "Công ty Cổ phần TECHPAL"),
        ("Địa chỉ", "D48-13 Khu đô thị Geleximco, đường Lê Trọng Tấn, phường Dương Nội, Hà Nội"),
        ("Thời gian", "20/07/2026 – 28/08/2026"),
        ("Người hướng dẫn", "Nguyễn Hồng Sơn"),
        ("Đại diện xác nhận", "Nguyễn Thế Tài, Phó Tổng giám đốc")
    ]
    for r_idx, (k, v) in enumerate(appA_data):
        c0 = t_appA.cell(r_idx, 0)
        c1 = t_appA.cell(r_idx, 1)
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.first_line_indent = Cm(0)
        r0 = p0.add_run(k)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(12)
        r0.bold = True
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.first_line_indent = Cm(0)
        r1 = p1.add_run(v)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)

    add_heading_2("A.2 Nội dung thực tập tóm tắt")
    add_para(
        "Tìm hiểu cơ cấu tổ chức và các lĩnh vực hoạt động của Công ty Cổ phần TECHPAL. "
        "Khảo sát danh mục linh kiện điện tử và thiết bị tự động hóa công nghiệp. "
        "Tìm hiểu phương pháp đọc mã linh kiện, datasheet và các tiêu chí kỹ thuật khi lựa chọn thiết bị. "
        "Phân tích vai trò của cảm biến, nguồn, rơ-le, PLC và mô-đun vào/ra trong hệ thống đo lường – điều khiển. "
        "Xây dựng yêu cầu và sơ đồ khối minh họa cho hệ thống giám sát nhiệt độ tủ kỹ thuật. "
        "Tổng hợp kết quả, nhận xét nhược điểm và đề xuất hướng hoàn thiện."
    )

    # ====================================================
    # PHỤ LỤC B: MINH CHỨNG THỰC TẬP
    # ====================================================
    doc.add_page_break()
    add_heading_1("PHỤ LỤC B: MINH CHỨNG THỰC TẬP")
    
    add_heading_2("B.1 Phiếu đánh giá quá trình thực tập")
    add_para(
        "Bản scan phiếu đánh giá do đơn vị tiếp nhận lập được trình bày trong Hình B.1 và Hình B.2."
    )
    
    # Page 1 scan
    p_sc1 = doc.add_paragraph()
    p_sc1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sc1.paragraph_format.space_before = Pt(6)
    p_sc1.paragraph_format.space_after = Pt(2)
    p_sc1.paragraph_format.first_line_indent = Cm(0)
    if os.path.exists("report/src/img/phieu-danh-gia-trang-1.jpg.png"):
        p_sc1.add_run().add_picture("report/src/img/phieu-danh-gia-trang-1.jpg.png", width=Cm(13.5))
    add_figure_caption("Hình B.1: Phiếu đánh giá quá trình thực tập (trang 1)")
    
    doc.add_page_break()
    # Page 2 scan
    p_sc2 = doc.add_paragraph()
    p_sc2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sc2.paragraph_format.space_before = Pt(6)
    p_sc2.paragraph_format.space_after = Pt(2)
    p_sc2.paragraph_format.first_line_indent = Cm(0)
    if os.path.exists("report/src/img/phieu-danh-gia-trang-2.jpg.png"):
        p_sc2.add_run().add_picture("report/src/img/phieu-danh-gia-trang-2.jpg.png", width=Cm(13.5))
    add_figure_caption("Hình B.2: Phiếu đánh giá quá trình thực tập (trang 2)")

    # Save document
    doc.save(output_path)
    print(f"Generated Word document successfully: {output_path}")

if __name__ == "__main__":
    os.makedirs("report", exist_ok=True)
    create_report_docx("report/report.docx")
    create_report_docx("report.docx")
